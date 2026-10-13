"""Shows what retrieval, the sufficiency judge and the web fallback do for one question.

Run from anywhere:
    python eval/probe_web_fallback.py --fixture progress --question "..."
    python eval/probe_web_fallback.py --corpus --question "..."
"""

from __future__ import annotations

import argparse
import asyncio
import io
import sys
import uuid
from contextlib import contextmanager
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import run_eval  # noqa: E402,F401  (sets cwd, sys.path and loads the env)

FIXTURES = {
    "progress": (
        "weekly-progress-report.docx",
        [
            ("Weekly Progress Report: Billing Data Migration", None),
            ("Summary", "The team moved the legacy invoice tables to the new warehouse this week. Overall status is green."),
            ("Completed this week", "Exported 4.2 million invoice rows. Wrote the mapping document for customer identifiers. Set up the staging database and the nightly sync job."),
            ("Risks", "The legacy export tool times out on tables larger than 2 million rows, so exports run in batches."),
            ("Plan for next week", "Build the schema validation module for the invoice tables. Run the first full reconciliation of row counts against the legacy system. Start load testing the staging database with 50 concurrent users."),
        ],
    ),
    "progress_labeled": (
        "weekly-progress-report-labeled.docx",
        [
            ("Weekly Progress Report: Billing Data Migration (Week 1)", None),
            ("Reporting period", "Week 1, the first week of the migration project."),
            ("Completed this week", "Exported 4.2 million invoice rows. Wrote the mapping document for customer identifiers. Set up the staging database and the nightly sync job."),
            ("Risks", "The legacy export tool times out on tables larger than 2 million rows, so exports run in batches."),
            ("Plan for next week", "Build the schema validation module for the invoice tables. Run the first full reconciliation of row counts against the legacy system. Start load testing the staging database with 50 concurrent users."),
        ],
    ),
    "progress_long": (
        "monthly-engineering-update.docx",
        [
            ("Engineering Update: Payments Platform Rewrite", None),
            ("Overview", "This update covers the first stretch of the payments platform rewrite. " + "The rewrite replaces the monolithic payment service with smaller services that own their data. " * 6),
            ("Architecture decisions", "The team chose an event log between services and a single shared gateway for external partners. " + "Each service exposes a versioned interface and publishes events after every committed change. " * 8),
            ("Work completed", "Delivered the ledger service skeleton, the idempotency layer and the first partner sandbox. " + "Code review turnaround averaged one day and test coverage on the ledger service reached 82 percent. " * 6),
            ("Incidents", "Two staging outages were traced to a misconfigured message broker. " + "Both were resolved within an hour and a runbook entry was added for each. " * 6),
            ("Look ahead", "Migrate the refund flow to the new ledger. Add reconciliation reports for finance. Begin the security review of the partner gateway. Prepare the cutover rehearsal."),
        ],
    ),
}


def build_docx(sections: list[tuple[str, str | None]]) -> bytes:
    from docx import Document

    doc = Document()
    for index, (heading, body) in enumerate(sections):
        if body is None:
            doc.add_heading(heading, level=0)
        else:
            doc.add_heading(heading, level=1)
            doc.add_paragraph(body)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@contextmanager
def no_stage(name, sub_question=None):
    yield


def show_passages(label: str, candidates, limit: int = 5) -> None:
    print(f"\n{label}")
    for i, cand in enumerate(candidates[:limit], start=1):
        dense = [p.score for p in cand.provenance if p.retriever == "vector" and p.score is not None]
        meta = cand.metadata
        where = f"{meta.get('source_doc')} p{meta.get('page')}" if cand.source == "corpus" else meta.get("source_url")
        best = f"{max(dense):.3f}" if dense else "n/a"
        print(f"  [{i}] {cand.source} {where} score={cand.score:.4f} dense={best}")
        print(f"      {' '.join(cand.text.split())[:170]}")


async def probe(question: str, fixture: str | None, answer: bool) -> None:
    from app.core.logging import configure_logging
    from app.ingestion.indexer import get_keyword_index, get_vector_store
    from app.ingestion.pipeline import ingest_files
    from app.query.pipeline import retrieve_and_rank
    from app.shared.session_store import PERSISTENT_SCOPE, get_session_stores, scope_for

    configure_logging()
    session_id = f"probe-{uuid.uuid4().hex[:12]}"
    if fixture:
        name, sections = FIXTURES[fixture]
        scope = scope_for(session_id)
        store, index = get_session_stores(session_id)
        result = ingest_files([(name, build_docx(sections))], scope, vector_store=store, keyword_index=index)
        print(f"uploaded {name}: {result.index.total_indexed} chunks into {scope}")
    else:
        scope, store, index = PERSISTENT_SCOPE, get_vector_store(), get_keyword_index()

    ranked = await retrieve_and_rank(
        question, no_stage, corpus_scope=scope, vector_store=store, keyword_index=index
    )
    print(f"\nQUESTION: {question}")
    show_passages("RETRIEVED (top fused, before web/rerank):", ranked.fused)
    s = ranked.sufficiency
    print(f"\nSUFFICIENCY: sufficient={s.sufficient} partial={s.partial} method={s.method} top_dense={s.scores.top_dense}")
    print(f"  reason: {s.reason}")
    print(f"  missing: {s.missing!r}")
    web = ranked.web_search
    if web is None:
        print("\nWEB SEARCH: not run")
    else:
        print(f"\nWEB SEARCH: status={web.status} failure={web.failure} kept={len(web.candidates)} dropped={web.dropped}")
        print(f"  query sent to Tavily: {web.query!r}")
        for c in web.candidates:
            print(f"  - {c.metadata.get('title')} | {c.metadata.get('source_url')}")
            print(f"      {c.text[:150]}")
    show_passages("FINAL CONTEXT (reranked, goes to the generator):", ranked.context.passages, limit=8)
    if answer:
        from app.query.pipeline import run_query

        result = await run_query(question, session_id, bypass_cache=True)
        kinds = sorted({c.kind for c in result.citations})
        print(f"\nANSWER ({result.terminated_at}, citation kinds={kinds}):\n{result.answer}")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--question", required=True)
    parser.add_argument("--answer", action="store_true", help="also run the full pipeline and print the answer")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--fixture", choices=sorted(FIXTURES))
    group.add_argument("--corpus", action="store_true")
    args = parser.parse_args()
    asyncio.run(probe(args.question, args.fixture, args.answer))
    return 0


if __name__ == "__main__":
    sys.exit(main())
