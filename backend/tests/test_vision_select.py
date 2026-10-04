"""Global, budgeted vision page selection. No network: everything here is local PDF analysis or a scripted model."""

import importlib.util
import io
import os
import sys
from pathlib import Path

import docx
import pymupdf
import pytest
from PIL import Image

from app.core.config import settings
from app.ingestion import vision_select as vs
from app.ingestion.extractors import vision
from app.ingestion.extractors.pdf import extract as pdf_extract
from app.ingestion.pipeline import ingest_files
from app.ingestion.vision_select import (
    PageFeatures,
    VisionCandidate,
    VisionPageCapExceeded,
    plan_vision,
    score_page,
    select,
    select_vision_pages,
)

CHART_REPLY = "CONTENT_TYPE: chart\nChart type: bar chart. Title: Quarterly revenue. Data points: Q1 10, Q2 12."
PROSE = "This page carries ordinary running prose about the annual report and nothing visual at all. " * 6


# --------------------------------------------------------------------------- #
# Builders
# --------------------------------------------------------------------------- #
def _noise_png(side: int = 400) -> bytes:
    buffer = io.BytesIO()
    Image.frombytes("RGB", (side, side), os.urandom(side * side * 3)).save(buffer, format="PNG")
    return buffer.getvalue()


def _text_page(doc, text: str = PROSE) -> None:
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 700), text, fontsize=10)


def _table_page(doc, rows: int = 6, cols: int = 4) -> None:
    page = doc.new_page()
    left, top, cell_w, cell_h = 50, 100, 100, 28
    for r in range(rows + 1):
        page.draw_line((left, top + r * cell_h), (left + cols * cell_w, top + r * cell_h))
    for c in range(cols + 1):
        page.draw_line((left + c * cell_w, top), (left + c * cell_w, top + rows * cell_h))
    for r in range(rows):
        for c in range(cols):
            page.insert_text((left + c * cell_w + 6, top + r * cell_h + 18), f"r{r}c{c}", fontsize=9)
    page.insert_text((50, 60), "Quarterly figures by region and product line for the year.", fontsize=10)


def _chart_page(doc, bars: int = 400) -> None:
    page = doc.new_page()
    for i in range(bars):  # one path per bar, as a charting library writes them
        x = 40 + (i % 100) * 5
        y = 100 + (i // 100) * 120
        page.draw_rect(pymupdf.Rect(x, y + 100 - (i % 37), x + 3, y + 100), fill=(0.2, 0.4, 0.8))
    page.insert_text((50, 60), "Figure 3. Monthly trend.", fontsize=10)


def _scanned_page(doc) -> None:
    page = doc.new_page()
    page.insert_image(page.rect, stream=_noise_png())


def _pdf(*builders) -> bytes:
    doc = pymupdf.open()
    for build in builders:
        build(doc)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_with_images(count: int) -> bytes:
    document = docx.Document()
    document.add_paragraph("Body text.")
    for _ in range(count):
        document.add_picture(io.BytesIO(_noise_png(96)))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _candidate(doc: str, unit: int, score: float, kind: str = "pdf_page") -> VisionCandidate:
    return VisionCandidate(doc, unit, kind, score, f"score {score}")


# --------------------------------------------------------------------------- #
# Scoring and page analysis
# --------------------------------------------------------------------------- #
def test_a_page_with_a_table_is_selected():
    selection = plan_vision([("t.pdf", _pdf(_table_page))])

    assert [(c.unit, c.kind) for c in selection.selected] == [(1, "pdf_page")]
    assert "table" in selection.selected[0].reason


def test_a_vector_chart_page_is_selected():
    selection = plan_vision([("c.pdf", _pdf(_chart_page))])

    assert [c.unit for c in selection.selected] == [1]
    assert "vector drawings" in selection.selected[0].reason


def test_a_plain_text_page_is_not_selected():
    selection = plan_vision([("p.pdf", _pdf(_text_page, _text_page))])

    assert selection.selected == [] and selection.scanned_to_ocr == []


def test_a_scanned_page_goes_to_ocr_and_does_not_use_the_vision_budget():
    selection = plan_vision([("s.pdf", _pdf(_scanned_page, _text_page, _scanned_page))])

    assert selection.selected == [] and selection.qualified == 0
    assert selection.scanned_to_ocr == [("s.pdf", 1), ("s.pdf", 3)]


def test_score_page_combines_signals_and_chart_dense_documents_win_ties():
    table_only = PageFeatures(text_chars=2000, image_coverage=0.0, drawings=10, tables=((8, 5),))
    table_and_chart = PageFeatures(text_chars=2000, image_coverage=0.0, drawings=900, tables=((8, 5),))

    plain, _ = score_page(table_only)
    dense, reason = score_page(table_only, chart_dense=True)
    both, _ = score_page(table_and_chart)

    assert plain == vs.TABLE_SCORE and dense == plain + vs.CHART_DENSE_BONUS and "chart-dense" in reason
    assert both == vs.TABLE_SCORE + vs.CHART_DRAWINGS_SCORE


def test_weak_signals_alone_do_not_qualify():
    small_table = PageFeatures(2000, 0.0, 10, ((3, 3),))
    few_drawings = PageFeatures(2000, 0.0, vs.MIN_CHART_DRAWINGS - 1)
    caption_beside_image = PageFeatures(60, 0.2, 0)

    assert all(score_page(f)[0] < vs.MIN_SELECT_SCORE for f in (small_table, few_drawings, caption_beside_image))
    assert score_page(caption_beside_image)[0] == vs.LOW_TEXT_IMAGE_SCORE


def test_a_large_image_on_a_page_with_text_is_a_strong_signal():
    page = PageFeatures(text_chars=900, image_coverage=0.7)

    assert score_page(page)[0] >= vs.MIN_SELECT_SCORE and not page.is_scanned


# --------------------------------------------------------------------------- #
# Global, budgeted selection
# --------------------------------------------------------------------------- #
def test_strict_mode_raises_when_the_plan_exceeds_the_cap():
    candidates = [_candidate(f"d{i}.pdf", 1, 4.0) for i in range(5)]

    with pytest.raises(VisionPageCapExceeded) as raised:
        select(candidates, cap=3, strict=True)

    assert raised.value.selected == 5 and raised.value.cap == 3 and "Tighten the thresholds" in str(raised.value)


def test_strict_mode_passes_when_the_plan_fits():
    candidates = [_candidate(f"d{i}.pdf", 1, 4.0) for i in range(3)]

    assert len(select(candidates, cap=3, strict=True).selected) == 3


def test_truncation_keeps_the_best_pages_spread_across_documents_and_records_the_rest():
    candidates = (
        [_candidate("big.pdf", n, 9.0 - n * 0.1) for n in range(1, 6)]
        + [_candidate("b.pdf", 1, 4.0), _candidate("c.pdf", 1, 3.5)]
    )

    selection = select(candidates, cap=3, strict=False, per_doc=10)

    assert [(c.source_doc, c.unit) for c in selection.selected] == [("b.pdf", 1), ("big.pdf", 1), ("c.pdf", 1)]
    assert len(selection.dropped) == 4 and all("page cap" in why for _, why in selection.dropped)


def test_one_document_cannot_take_more_than_the_per_document_ceiling():
    candidates = [_candidate("big.pdf", n, 5.0) for n in range(1, 12)] + [_candidate("small.pdf", 1, 3.5)]

    selection = select(candidates, cap=80, strict=True)

    assert sum(c.source_doc == "big.pdf" for c in selection.selected) == vs.VISION_MAX_PAGES_PER_DOC
    assert any(c.source_doc == "small.pdf" for c in selection.selected)
    assert len(selection.dropped) == 11 - vs.VISION_MAX_PAGES_PER_DOC


def test_docx_and_standalone_images_count_against_the_same_cap(monkeypatch):
    files = [
        ("table.pdf", _pdf(_table_page)),
        ("report.docx", _docx_with_images(2)),
        ("chart.png", _noise_png()),
    ]

    selection = plan_vision(files)
    kinds = sorted(c.kind for c in selection.selected)

    assert kinds == ["docx_image", "docx_image", "image", "pdf_page"]
    with pytest.raises(VisionPageCapExceeded):
        plan_vision(files, cap=3, strict=True)
    assert len(plan_vision(files, cap=3, strict=False).selected) == 3


def test_standalone_images_outrank_page_guesses_when_the_cap_bites():
    files = [("table.pdf", _pdf(_table_page)), ("chart.png", _noise_png())]

    selection = plan_vision(files, cap=1, strict=False)

    assert [c.kind for c in selection.selected] == ["image"]


def test_selection_is_deterministic_and_independent_of_input_order():
    files = [
        ("a.pdf", _pdf(_table_page, _text_page, _chart_page)),
        ("b.docx", _docx_with_images(1)),
        ("c.png", _noise_png()),
    ]

    first = plan_vision(files, cap=3, strict=False)
    second = plan_vision(list(reversed(files)), cap=3, strict=False)

    assert [(c.source_doc, c.unit, c.score) for c in first.selected] == [
        (c.source_doc, c.unit, c.score) for c in second.selected
    ]
    assert [(c.source_doc, c.unit) for c, _ in first.dropped] == [(c.source_doc, c.unit) for c, _ in second.dropped]


def test_select_vision_pages_dry_run_returns_tuples_and_makes_no_call(monkeypatch, tmp_path):
    def boom(*args, **kwargs):
        raise AssertionError("the dry run called the model")

    monkeypatch.setattr(vision._client, "generate_vision", boom)
    (tmp_path / "t.pdf").write_bytes(_pdf(_table_page, _text_page))
    (tmp_path / "i.png").write_bytes(_noise_png())

    rows = select_vision_pages(sorted(tmp_path.iterdir()), chart_dense_docs=frozenset({"t.pdf"}))

    assert [(doc, label) for doc, label, _, _ in rows] == [("i.png", "image"), ("t.pdf", 1)]
    assert rows[0][2] == vs.STANDALONE_IMAGE_SCORE and rows[1][2] == vs.TABLE_SCORE + vs.CHART_DENSE_BONUS
    assert "table" in rows[1][3]


# --------------------------------------------------------------------------- #
# The plan drives extraction
# --------------------------------------------------------------------------- #
@pytest.fixture
def scripted(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "VISION_CACHE_DIR", str(tmp_path / "vision-cache"))
    calls = {"vision": 0, "ocr": 0}

    def generate_vision(*args, **kwargs):
        calls["vision"] += 1
        return CHART_REPLY

    def run_ocr(image):
        calls["ocr"] += 1
        return "text read by ocr from the scanned page"

    monkeypatch.setattr(vision._client, "generate_vision", generate_vision)
    monkeypatch.setattr(vision, "run_ocr", run_ocr)
    return calls


def test_only_planned_pages_reach_vision_and_their_text_layer_is_kept(scripted):
    content = _pdf(_text_page, _table_page, _text_page)

    pieces = pdf_extract(content, "doc.pdf", vision_units={2})

    assert scripted == {"vision": 1, "ocr": 0}
    by_method = sorted((p["page"], p["extraction_method"]) for p in pieces)
    assert by_method == [(1, "text"), (2, "text"), (2, "vision"), (3, "text")]


def test_a_scanned_page_is_read_by_ocr_without_a_vision_call(scripted):
    pieces = pdf_extract(_pdf(_scanned_page, _text_page), "doc.pdf", vision_units=set())

    assert scripted == {"vision": 0, "ocr": 1}
    assert [(p["page"], p["extraction_method"]) for p in pieces] == [(1, "ocr"), (2, "text")]


def test_a_pdf_extracted_on_its_own_plans_for_itself(scripted):
    pieces = pdf_extract(_pdf(_text_page, _table_page), "doc.pdf")

    assert scripted["vision"] == 1 and {p["extraction_method"] for p in pieces if p["page"] == 2} == {"text", "vision"}


def test_a_batch_over_the_cap_raises_instead_of_sending_pages_to_ocr(scripted, monkeypatch):
    monkeypatch.setattr(settings, "MAX_VISION_PAGES", 2)
    items = [vision.VisionPage(image_bytes=_noise_png(32), page=n) for n in range(3)]

    with pytest.raises(VisionPageCapExceeded):
        vision.extract_pages(items)

    assert scripted == {"vision": 0, "ocr": 0}


def test_ingest_in_strict_mode_raises_before_any_model_call(scripted, monkeypatch):
    monkeypatch.setattr(settings, "MAX_VISION_PAGES", 1)
    files = [("a.pdf", _pdf(_table_page)), ("b.pdf", _pdf(_chart_page))]

    with pytest.raises(VisionPageCapExceeded):
        ingest_files(files, corpus_scope="persistent")

    assert scripted == {"vision": 0, "ocr": 0}


def test_ingest_with_truncation_keeps_the_best_and_records_what_it_dropped(scripted, monkeypatch):
    monkeypatch.setattr(settings, "MAX_VISION_PAGES", 1)
    files = [("a.pdf", _pdf(_table_page)), ("b.png", _noise_png())]

    result = ingest_files(files, corpus_scope="persistent", vision_strict=False)

    assert scripted["vision"] == 1
    assert [c.source_doc for c in result.vision.selected] == ["b.png"]
    assert [(c.source_doc, c.unit) for c, _ in result.vision.dropped] == [("a.pdf", 1)]
    assert {c["extraction_method"] for c in result.chunks if c["source_doc"] == "a.pdf"} == {"text"}


def test_ingest_extracts_planned_tables_and_charts_through_vision(scripted):
    result = ingest_files([("a.pdf", _pdf(_text_page, _chart_page))], corpus_scope="persistent")

    methods = {(c["page"], c["extraction_method"], c["chunk_type"]) for c in result.chunks}
    assert (2, "vision", "chart") in methods and (1, "vision", "chart") not in methods
    assert result.vision.scanned_to_ocr == [] and len(result.vision.selected) == 1


# --------------------------------------------------------------------------- #
# The ingestion script
# --------------------------------------------------------------------------- #
def _load_script():
    path = Path(__file__).resolve().parents[1] / "scripts" / "ingest_corpus.py"
    spec = importlib.util.spec_from_file_location("ingest_corpus_script", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_plan_vision_flag_prints_the_selection_and_exits_without_ingesting(tmp_path, monkeypatch, capsys):
    (tmp_path / "t.pdf").write_bytes(_pdf(_table_page, _text_page))
    (tmp_path / "i.png").write_bytes(_noise_png())
    script = _load_script()
    monkeypatch.setattr(sys, "argv", ["ingest_corpus.py", "--plan-vision", "--corpus-dir", str(tmp_path)])
    monkeypatch.setattr(script, "ingest_files", lambda *a, **k: pytest.fail("ingested during --plan-vision"))

    code = script.main()
    out = capsys.readouterr().out

    assert code == 0 and "Vision plan: 2 of" in out and "t.pdf\t1\t" in out and "i.png\timage\t" in out
    assert "documents with at least one selected page: 2 of 2" in out


def test_the_script_is_strict_by_default_and_truncates_only_when_asked(tmp_path, monkeypatch, capsys):
    (tmp_path / "t.pdf").write_bytes(_pdf(_table_page))
    (tmp_path / "i.png").write_bytes(_noise_png())
    script = _load_script()
    monkeypatch.setattr(settings, "MAX_VISION_PAGES", 1)
    seen = {}

    def fake_ingest(files, corpus_scope, **kwargs):
        seen.update(kwargs)
        raise VisionPageCapExceeded(2, 1) if kwargs["vision_strict"] else SystemExit(0)

    monkeypatch.setattr(script, "ingest_files", fake_ingest)
    monkeypatch.setattr(sys, "argv", ["ingest_corpus.py", "--corpus-dir", str(tmp_path)])
    assert script.main() == 2 and seen["vision_strict"] is True
    assert "--allow-truncation" in capsys.readouterr().err

    monkeypatch.setattr(sys, "argv", ["ingest_corpus.py", "--corpus-dir", str(tmp_path), "--allow-truncation"])
    with pytest.raises(SystemExit):
        script.main()
    assert seen["vision_strict"] is False


# --------------------------------------------------------------------------- #
# Fallback: one best page for a chart-dense document with no normally scored page
# --------------------------------------------------------------------------- #
def _grid_page(doc, rows: int, cols: int) -> None:
    page = doc.new_page()
    left, top, cell_w, cell_h = 50, 100, 100, 28
    for r in range(rows + 1):
        page.draw_line((left, top + r * cell_h), (left + cols * cell_w, top + r * cell_h))
    for c in range(cols + 1):
        page.draw_line((left + c * cell_w, top), (left + c * cell_w, top + rows * cell_h))
    for r in range(rows):
        for c in range(cols):
            page.insert_text((left + c * cell_w + 6, top + r * cell_h + 18), f"r{r}c{c}", fontsize=9)
    page.insert_textbox(pymupdf.Rect(50, 300, 550, 700), PROSE, fontsize=10)


def _image_beside_text_page(doc) -> None:
    """Text plus an image covering about 38% of the page: a relaxed signal, not a strong one."""
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(30, 60, 565, 420), stream=_noise_png())
    page.insert_textbox(pymupdf.Rect(50, 440, 550, 800), PROSE, fontsize=10)


def _drawings_page(doc) -> None:
    """150 separate paths: above the fallback bar of 100, below the chart bar of 300."""
    page = doc.new_page()
    for i in range(150):
        page.draw_rect(pymupdf.Rect(40 + (i % 50) * 10, 100 + (i // 50) * 60, 46 + (i % 50) * 10, 150 + (i // 50) * 60), fill=(0.5, 0.5, 0.5))
    page.insert_textbox(pymupdf.Rect(50, 300, 550, 700), PROSE, fontsize=10)


DENSE = frozenset({"doc.pdf"})


def test_a_chart_dense_document_with_no_strong_page_gets_exactly_one_fallback_page():
    content = _pdf(_text_page, _image_beside_text_page, _text_page)

    selection = plan_vision([("doc.pdf", content)], chart_dense_docs=DENSE)

    assert [(c.unit, c.reason, c.fallback) for c in selection.selected] == [(2, "fallback_best_page", True)]
    assert selection.selected[0].score == vs.FALLBACK_SCORE < vs.MIN_SELECT_SCORE


def test_a_document_with_a_strong_page_gets_no_fallback_page():
    content = _pdf(_table_page, _image_beside_text_page)

    selection = plan_vision([("doc.pdf", content)], chart_dense_docs=DENSE)

    assert [(c.unit, c.fallback) for c in selection.selected] == [(1, False)]


def test_a_2x2_table_alone_does_not_qualify_but_a_3x3_table_does():
    two = plan_vision([("doc.pdf", _pdf(_text_page, lambda d: _grid_page(d, 2, 2)))], chart_dense_docs=DENSE)
    three = plan_vision([("doc.pdf", _pdf(_text_page, lambda d: _grid_page(d, 3, 3)))], chart_dense_docs=DENSE)

    assert two.selected == [] and "no page has a strong or relaxed visual signal" in two.notes["doc.pdf"]
    assert [(c.unit, c.fallback) for c in three.selected] == [(2, True)]


def test_enough_vector_drawings_qualify_for_the_fallback():
    selection = plan_vision([("doc.pdf", _pdf(_text_page, _drawings_page))], chart_dense_docs=DENSE)

    assert [(c.unit, c.fallback) for c in selection.selected] == [(2, True)]


def test_documents_that_are_not_chart_dense_never_get_a_fallback_page():
    content = _pdf(_text_page, _image_beside_text_page)

    selection = plan_vision([("doc.pdf", content)], chart_dense_docs=frozenset())

    assert selection.selected == []


def test_at_most_one_fallback_page_per_document():
    content = _pdf(_image_beside_text_page, _image_beside_text_page, _drawings_page)

    selection = plan_vision([("doc.pdf", content)], chart_dense_docs=DENSE)

    assert len(selection.selected) == 1 and selection.selected[0].fallback


def test_the_fallback_page_is_the_best_by_the_tie_break_order():
    content = _pdf(_image_beside_text_page, _drawings_page)  # more drawings beats a larger image

    selection = plan_vision([("doc.pdf", content)], chart_dense_docs=DENSE)

    assert [c.unit for c in selection.selected] == [2]


def test_a_scanned_only_document_still_goes_to_ocr_with_no_fallback():
    selection = plan_vision([("doc.pdf", _pdf(_scanned_page, _scanned_page))], chart_dense_docs=DENSE)

    assert selection.selected == [] and selection.scanned_to_ocr == [("doc.pdf", 1), ("doc.pdf", 2)]
    assert "scanned" in selection.notes["doc.pdf"]


def _fallback(doc: str, unit: int = 1) -> VisionCandidate:
    return VisionCandidate(doc, unit, "pdf_page", vs.FALLBACK_SCORE, "fallback_best_page", fallback=True)


def test_fallback_pages_are_dropped_first_under_truncation():
    normal = [_candidate(f"n{i}.pdf", 1, 3.5) for i in range(3)]
    extras = [_fallback("f1.pdf"), _fallback("f2.pdf")]

    exact = select(normal + extras, cap=3, strict=False)
    one_spare = select(normal + extras, cap=4, strict=False)

    assert [c.source_doc for c in exact.selected] == ["n0.pdf", "n1.pdf", "n2.pdf"]
    assert [c.source_doc for c, _ in exact.dropped] == ["f1.pdf", "f2.pdf"]
    assert [c.source_doc for c in one_spare.selected] == ["f1.pdf", "n0.pdf", "n1.pdf", "n2.pdf"]
    assert [c.source_doc for c, _ in one_spare.dropped] == ["f2.pdf"]


def test_a_fallback_page_beats_nothing_but_never_a_second_page_of_a_normal_document():
    candidates = [_candidate("a.pdf", 1, 5.0), _candidate("a.pdf", 2, 4.0), _fallback("f.pdf")]

    selection = select(candidates, cap=2, strict=False)

    assert [(c.source_doc, c.unit) for c in selection.selected] == [("a.pdf", 1), ("a.pdf", 2)]


def test_fallback_pages_count_toward_the_strict_cap():
    candidates = [_candidate(f"n{i}.pdf", 1, 3.5) for i in range(3)] + [_fallback("f.pdf")]

    with pytest.raises(VisionPageCapExceeded) as raised:
        select(candidates, cap=3, strict=True)

    assert raised.value.selected == 4
    assert len(select(candidates, cap=4, strict=True).selected) == 4


def test_selection_with_fallback_pages_is_deterministic():
    files = [
        ("a.pdf", _pdf(_text_page, _image_beside_text_page)),
        ("b.pdf", _pdf(_text_page, _drawings_page)),
        ("c.pdf", _pdf(_table_page)),
    ]
    dense = frozenset({"a.pdf", "b.pdf", "c.pdf"})

    first = plan_vision(files, chart_dense_docs=dense)
    second = plan_vision(list(reversed(files)), chart_dense_docs=dense)

    assert [(c.source_doc, c.unit, c.score, c.fallback) for c in first.selected] == [
        (c.source_doc, c.unit, c.score, c.fallback) for c in second.selected
    ]
    assert [(c.source_doc, c.fallback) for c in first.selected] == [("a.pdf", True), ("b.pdf", True), ("c.pdf", False)]
