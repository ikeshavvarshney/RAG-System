"""DOCX sectioning, tables as markdown, and embedded images (D-18). No network: vision and OCR are scripted."""

import io
import os

import docx
import pytest
from PIL import Image

from app.core.config import settings
from app.ingestion.extractors import docx as docx_extractor
from app.ingestion.extractors import vision
from app.ingestion.extractors.docx import (
    MIN_IMAGE_BYTES,
    MIN_IMAGE_SIDE_PX,
    images_for_vision,
    plan_embedded_images,
)
from app.ingestion.extractors.ocr import OCRUnavailable
from app.ingestion.pipeline import ingest_files
from app.ingestion.splitter import _token_len

CHART_REPLY = "CONTENT_TYPE: chart\nChart type: bar chart. Title: Quarterly revenue. Data points: Q1 10, Q2 12."


def _noise_png(side: int = 96) -> bytes:
    """An incompressible image, so its byte size clears the decorative threshold."""
    buffer = io.BytesIO()
    Image.frombytes("RGB", (side, side), os.urandom(side * side * 3)).save(buffer, format="PNG")
    return buffer.getvalue()


def _words(count: int, tag: str) -> str:
    """About ``count`` tokens of plain prose that starts with a unique, easy-to-find marker word."""
    return f"{tag} " + " ".join(["revenue"] * (count - 1))


def _docx(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _picture(document, data: bytes) -> None:
    document.add_picture(io.BytesIO(data))


@pytest.fixture
def vision_mocks(monkeypatch, tmp_path):
    """A scripted vision model and OCR, an empty vision cache, and a count of every call."""
    monkeypatch.setattr(settings, "VISION_CACHE_DIR", str(tmp_path / "vision-cache"))
    calls = {"vision": 0, "ocr": 0}
    state = {"vision_reply": CHART_REPLY, "ocr_reply": "scanned text from the image", "ocr_error": None}

    def generate_vision(*args, **kwargs):
        calls["vision"] += 1
        reply = state["vision_reply"]
        if isinstance(reply, Exception):
            raise reply
        return reply

    def run_ocr(image):
        calls["ocr"] += 1
        if state["ocr_error"] is not None:
            raise state["ocr_error"]
        return state["ocr_reply"]

    monkeypatch.setattr(vision._client, "generate_vision", generate_vision)
    monkeypatch.setattr(vision, "run_ocr", run_ocr)
    return calls, state


def test_paragraphs_under_one_heading_merge_into_one_chunk_in_range():
    def build(d):
        d.add_heading("Methods", level=1)
        for i in range(7):
            d.add_paragraph(_words(45, f"pmark{i}"))

    chunks = ingest_files([("doc.docx", _docx(build))], corpus_scope="persistent").chunks

    assert len(chunks) == 1
    assert chunks[0]["location"] == "Methods" and chunks[0]["page"] is None
    assert settings.CHUNK_MIN_TOKENS <= _token_len(chunks[0]["text"]) <= settings.CHUNK_MAX_TOKENS
    assert all(f"pmark{i}" in chunks[0]["text"] for i in range(7))


def test_each_section_is_cut_at_its_heading_and_carries_it_in_location():
    def build(d):
        d.add_heading("Background", level=1)
        for i in range(7):
            d.add_paragraph(_words(45, f"amark{i}"))
        d.add_heading("Results", level=1)
        for i in range(7):
            d.add_paragraph(_words(45, f"bmark{i}"))

    chunks = ingest_files([("doc.docx", _docx(build))], corpus_scope="persistent").chunks

    assert [c["location"] for c in chunks] == ["Background", "Results"]
    assert "bmark0" not in chunks[0]["text"] and "amark0" not in chunks[1]["text"]


def test_a_table_becomes_one_markdown_chunk_not_one_per_cell():
    def build(d):
        d.add_paragraph("Quarterly figures are shown below for every region in the report.")
        table = d.add_table(rows=4, cols=3)
        for r in range(4):
            for c in range(3):
                table.cell(r, c).text = f"r{r}c{c}"

    chunks = ingest_files([("doc.docx", _docx(build))], corpus_scope="persistent").chunks
    tables = [c for c in chunks if c["chunk_type"] == "table"]

    assert len(tables) == 1 and len(chunks) == 2
    lines = tables[0]["text"].splitlines()
    assert lines[0] == "| r0c0 | r0c1 | r0c2 |" and lines[1] == "| --- | --- | --- |" and len(lines) == 5


def test_merged_cells_and_pipes_keep_the_table_well_formed():
    def build(d):
        table = d.add_table(rows=2, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Header | merged"
        table.cell(0, 2).text = "Last"
        for c in range(3):
            table.cell(1, c).text = f"v{c}"

    lines = docx_extractor.extract(_docx(build), "t.docx")[0]["text"].splitlines()

    assert lines[0] == "| Header \\| merged |  | Last |"  # the merged cell spans two columns
    assert all(line.startswith("|") and line.endswith("|") for line in lines)


def test_a_genuinely_tiny_document_stays_one_small_chunk():
    chunks = ingest_files(
        [("doc.docx", _docx(lambda d: d.add_paragraph("A short note with a handful of words only.")))],
        corpus_scope="persistent",
    ).chunks

    assert len(chunks) == 1 and _token_len(chunks[0]["text"]) < settings.CHUNK_MIN_TOKENS


def test_embedded_image_with_a_vision_answer_is_a_vision_chunk(vision_mocks):
    calls, _ = vision_mocks
    content = _docx(lambda d: (d.add_paragraph("Figure follows."), _picture(d, _noise_png())))

    chunks = ingest_files([("doc.docx", content)], corpus_scope="persistent").chunks
    image = [c for c in chunks if c["location"] == "embedded_image_1"]

    assert len(image) == 1 and calls == {"vision": 1, "ocr": 0}
    assert image[0]["extraction_method"] == "vision" and image[0]["chunk_type"] == "chart"
    assert image[0]["source_doc"] == "doc.docx" and image[0]["page"] is None


def test_embedded_image_falls_back_to_ocr_and_is_tagged_ocr(vision_mocks):
    calls, state = vision_mocks
    state["vision_reply"] = RuntimeError("vision is down")
    content = _docx(lambda d: _picture(d, _noise_png()))

    chunks = ingest_files([("doc.docx", content)], corpus_scope="persistent").chunks

    assert calls == {"vision": 1, "ocr": 1} and len(chunks) == 1
    assert chunks[0]["extraction_method"] == "ocr" and chunks[0]["chunk_type"] == "text"
    assert chunks[0]["location"] == "embedded_image_1" and chunks[0]["page"] is None


def test_decorative_images_are_skipped_without_any_model_call(vision_mocks):
    calls, _ = vision_mocks
    tiny = io.BytesIO()
    Image.new("RGB", (16, 16), "red").save(tiny, format="PNG")
    narrow = _noise_png(side=MIN_IMAGE_SIDE_PX - 24)
    assert len(tiny.getvalue()) < MIN_IMAGE_BYTES <= len(narrow)
    content = _docx(lambda d: (d.add_paragraph("Body."), _picture(d, tiny.getvalue()), _picture(d, narrow)))

    plan = plan_embedded_images(content)
    pieces = docx_extractor.extract(content, "doc.docx")

    assert plan.kept == [] and [n for n, _ in plan.skipped] == [1, 2]
    assert "bytes" in plan.skipped[0][1] and "px" in plan.skipped[1][1]
    assert calls == {"vision": 0, "ocr": 0} and [p["extraction_method"] for p in pieces] == ["text"]


def test_images_for_vision_is_a_dry_run_bounded_by_the_vision_budget(vision_mocks, monkeypatch):
    calls, _ = vision_mocks
    monkeypatch.setattr(settings, "MAX_VISION_PAGES", 2)
    content = _docx(lambda d: [_picture(d, _noise_png()) for _ in range(3)])

    selected = images_for_vision(content)

    assert [i.location for i in selected] == ["embedded_image_1", "embedded_image_2"]
    assert len(plan_embedded_images(content).kept) == 3
    assert all(i.mime_type == "image/png" and i.width == i.height == 96 for i in selected)
    assert calls == {"vision": 0, "ocr": 0}


def test_images_beyond_the_vision_budget_go_to_ocr(vision_mocks, monkeypatch):
    calls, _ = vision_mocks
    monkeypatch.setattr(settings, "MAX_VISION_PAGES", 1)
    content = _docx(lambda d: [_picture(d, _noise_png()) for _ in range(2)])

    pieces = docx_extractor.extract(content, "doc.docx")

    assert [(p["location"], p["extraction_method"]) for p in pieces] == [
        ("embedded_image_1", "vision"),
        ("embedded_image_2", "ocr"),
    ]
    assert calls == {"vision": 1, "ocr": 1}


def test_an_image_that_fails_vision_and_ocr_is_dropped_but_the_text_survives(vision_mocks):
    _, state = vision_mocks
    state["vision_reply"] = RuntimeError("vision is down")
    state["ocr_error"] = OCRUnavailable("no engine")
    content = _docx(lambda d: (d.add_paragraph("The surviving body text of the document."), _picture(d, _noise_png())))

    result = ingest_files([("doc.docx", content)], corpus_scope="persistent")

    assert result.succeeded == ["doc.docx"] and result.failed == []
    assert [c["text"] for c in result.chunks] == ["The surviving body text of the document."]


def test_a_non_native_image_format_is_reencoded_to_png(vision_mocks):
    buffer = io.BytesIO()
    Image.frombytes("RGB", (96, 96), os.urandom(96 * 96 * 3)).save(buffer, format="BMP")
    content = _docx(lambda d: _picture(d, buffer.getvalue()))

    selected = images_for_vision(content)

    assert len(selected) == 1 and selected[0].mime_type == "image/png"
    assert selected[0].data.startswith(b"\x89PNG")
